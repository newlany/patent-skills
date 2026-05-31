from __future__ import annotations

import json
import os
import re
import subprocess
import time
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable, Iterator, Sequence
from urllib.parse import quote_plus

import pdfplumber
import requests
from bs4 import BeautifulSoup
from docx import Document
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/137.0.0.0 Safari/537.36"
)

SUPPORTED_TEXT_EXTENSIONS = {".docx", ".pdf", ".md", ".txt"}
COMPARISON_DIRNAME = "comparison"
EVIDENCE_DIRNAME = "evidence"
OTHER_DIRNAME = "other"

CITATION_KEYWORDS = [
    "对比文件",
    "对比文献",
    "比对文件",
    "对照文件",
    "引证文献",
    "引用文献",
    "现有技术文献",
    "专利文献",
    "对比专利",
    "证据",
    "附件证据",
    "附件",
    "prior art",
    "reference",
    "references",
    "evidence",
]

LABEL_PATTERN = re.compile(
    r"(?:(?:对比文件|对比文献|比对文件|对照文件|引证文献|现有技术文献|专利文献|证据|附件证据|附件)\s*[A-Z]?\d{1,3}|(?<![A-Z0-9])[DE]\d{1,3}(?![A-Z0-9]))",
    re.IGNORECASE,
)

PUBLICATION_PATTERN = re.compile(
    r"\b(?:CN|US|EP|WO|JP|KR|TW|DE|FR|GB|CA|AU|RU|BR|IN|MX|ES|IT|PT|AT|BE|CH|DK|FI|NO|SE|NL|SG|MY|VN|HK|EA|EM|GC|IB)"
    r"[A-Z0-9./-]{5,24}[A-Z]\d?\b",
    re.IGNORECASE,
)

APPLICATION_PATTERNS = [
    re.compile(r"\bPCT/[A-Z]{2}\d{4}/\d{4,7}\b", re.IGNORECASE),
    re.compile(r"\b(?:CN|US|EP|WO|JP|KR|TW|DE|FR|GB|CA|AU|RU|BR|IN|MX|ES|IT|PT|AT|BE|CH|DK|FI|NO|SE|NL|SG|MY|VN|HK|EA|EM|GC|IB)\d{8,14}(?:\.\d)?\b", re.IGNORECASE),
    re.compile(r"\b\d{8,14}(?:\.\d)?\b"),
]

TITLE_QUOTE_PATTERN = re.compile(r"[“\"]([^”\"\n]{4,180})[”\"]")
TITLE_AFTER_NUMBER_PATTERN = re.compile(
    r"(?:公开号|申请号|专利号|Publication(?: Number)?|Application(?: Number)?)"
    r"\s*[:：]?\s*[A-Z0-9./-]+\s*[，,:：]?\s*([^；;。\n]{4,180})",
    re.IGNORECASE,
)
COMPARISON_LABEL_INDEX_PATTERN = re.compile(r"(?:对比文件|D)\s*([0-9]{1,3})", re.IGNORECASE)
COMPARISON_SECTION_MARKER = "本通知书引用下列对比文件"
COMPARISON_SECTION_END_MARKERS = [
    "审查的结论性意见",
    "关于说明书",
    "关于权利要求书",
]
APPLICATION_NUMBER_FIELD_PATTERN = re.compile(
    r"申请号\s*[:：]\s*([A-Z]{0,2}\d{8,14}(?:\.\d)?)",
    re.IGNORECASE,
)
PUBLICATION_NUMBER_FIELD_PATTERN = re.compile(
    r"(?:公开号|公开(?:专利)?号)\s*[:：]\s*([A-Z]{2}[A-Z0-9./-]{5,24}[A-Z]\d?)",
    re.IGNORECASE,
)
INVENTION_TITLE_FIELD_PATTERN = re.compile(
    r"(?:发明创造名称|发明名称)\s*[:：]\s*(.+)",
    re.IGNORECASE,
)
SOURCE_APPLICATION_LABEL = "原申请文件"


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_space(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize_identifier(value: str) -> str:
    return re.sub(r"\s+", "", value or "").upper()


def normalize_title(value: str) -> str:
    value = normalize_space(value)
    return re.sub(r"[\W_]+", "", value).lower()


def unique_preserve(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for item in items:
        if not item:
            continue
        if item in seen:
            continue
        seen.add(item)
        ordered.append(item)
    return ordered


def sanitize_filename(value: str, max_length: int = 180) -> str:
    value = normalize_space(value)
    value = re.sub(r'[<>:"/\\|?*]+', "_", value)
    value = value.rstrip(". ")
    if len(value) > max_length:
        value = value[:max_length].rstrip(" ._")
    return value or "download"


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="ignore")


def read_docx_text(path: Path) -> str:
    document = Document(path)
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_space(paragraph.text)
        if text:
            chunks.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [normalize_space(cell.text) for cell in row.cells if normalize_space(cell.text)]
            if values:
                chunks.append(" | ".join(values))
    return "\n".join(chunks)


def read_pdf_text(path: Path) -> str:
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            text = normalize_space(text)
            if text:
                parts.append(text)
    return "\n".join(parts)


def read_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".pdf":
        return read_pdf_text(path)
    if suffix in {".md", ".txt"}:
        return read_text_file(path)
    raise ValueError(f"Unsupported file type: {path}")


def iter_supported_files(inputs: Sequence[Path]) -> Iterator[Path]:
    for item in inputs:
        if item.is_dir():
            for child in sorted(item.rglob("*")):
                if child.is_file() and child.suffix.lower() in SUPPORTED_TEXT_EXTENSIONS:
                    yield child
            continue
        if item.is_file() and item.suffix.lower() in SUPPORTED_TEXT_EXTENSIONS:
            yield item


def resolve_input_files(raw_inputs: Sequence[str | os.PathLike[str]]) -> list[Path]:
    paths = [Path(item).expanduser().resolve() for item in raw_inputs]
    files = unique_preserve(str(path) for path in iter_supported_files(paths))
    return [Path(path) for path in files]


def looks_like_citation_context(context: str) -> bool:
    lowered = context.lower()
    if LABEL_PATTERN.search(context):
        return True
    return any(keyword.lower() in lowered for keyword in CITATION_KEYWORDS)


def infer_category(label: str | None, context: str) -> str:
    label_upper = (label or "").upper()
    if label_upper.startswith("D") or "对比" in context or "引证" in context or "现有技术" in context:
        return "comparison"
    if label_upper.startswith("E") or "证据" in context:
        return "evidence"
    return "other"


def extract_labels(text: str) -> list[str]:
    cleaned = [normalize_space(match.group(0)) for match in LABEL_PATTERN.finditer(text)]
    return unique_preserve(cleaned)


def extract_publication_numbers(text: str) -> list[str]:
    values = [normalize_identifier(match.group(0)) for match in PUBLICATION_PATTERN.finditer(text)]
    return unique_preserve(values)


def is_likely_application_number(value: str) -> bool:
    normalized = normalize_identifier(value)
    if PUBLICATION_PATTERN.fullmatch(normalized):
        return False
    if normalized.startswith("PCT/"):
        return True
    digits = re.sub(r"\D", "", normalized)
    return 8 <= len(digits) <= 14


def extract_application_numbers(text: str) -> list[str]:
    values: list[str] = []
    for pattern in APPLICATION_PATTERNS:
        for match in pattern.finditer(text):
            raw = normalize_identifier(match.group(0))
            if is_likely_application_number(raw):
                values.append(raw)
    return unique_preserve(values)


def extract_title(text: str) -> str | None:
    quote_match = TITLE_QUOTE_PATTERN.search(text)
    if quote_match:
        return normalize_space(quote_match.group(1))
    after_number = TITLE_AFTER_NUMBER_PATTERN.search(text)
    if after_number:
        return normalize_space(after_number.group(1))
    lines = [normalize_space(line) for line in text.splitlines() if normalize_space(line)]
    for line in lines:
        if PUBLICATION_PATTERN.search(line) or any(keyword in line for keyword in CITATION_KEYWORDS):
            tail = PUBLICATION_PATTERN.sub("", line)
            tail = re.sub(r"\b(?:CN|US|EP|WO|JP|KR|TW|DE|FR|GB|CA|AU|RU|BR|IN|MX|ES|IT|PT|AT|BE|CH|DK|FI|NO|SE|NL|SG|MY|VN|HK|EA|EM|GC|IB)\d+(?:\.\d+)?\b", "", tail, flags=re.IGNORECASE)
            tail = tail.strip(" ：:;；,，")
            if 4 <= len(tail) <= 180:
                return tail
    return None


@dataclass
class ExtractedReference:
    label: str | None
    category: str
    publication_numbers: list[str] = field(default_factory=list)
    application_numbers: list[str] = field(default_factory=list)
    title: str | None = None
    context: str | None = None
    source_files: list[str] = field(default_factory=list)
    query_terms: list[str] = field(default_factory=list)

    def merge(self, other: "ExtractedReference") -> None:
        self.publication_numbers = unique_preserve(self.publication_numbers + other.publication_numbers)
        self.application_numbers = unique_preserve(self.application_numbers + other.application_numbers)
        if not self.title and other.title:
            self.title = other.title
        if not self.context and other.context:
            self.context = other.context
        self.source_files = unique_preserve(self.source_files + other.source_files)
        self.query_terms = unique_preserve(self.query_terms + other.query_terms)
        if self.label is None:
            self.label = other.label
        if self.category == "other" and other.category != "other":
            self.category = other.category

    def to_dict(self) -> dict:
        return asdict(self)


def candidate_key(reference: ExtractedReference) -> str:
    publication_key = "|".join(reference.publication_numbers)
    application_key = "|".join(reference.application_numbers)
    title_key = normalize_title(reference.title or "")
    label_key = normalize_identifier(reference.label or "")
    return "::".join([publication_key, application_key, title_key, label_key, reference.category])


def comparison_index_from_label(label: str | None) -> int | None:
    if not label:
        return None
    match = COMPARISON_LABEL_INDEX_PATTERN.search(label)
    if not match:
        return None
    return int(match.group(1))


def canonical_comparison_label(label: str | None, fallback_index: int | None = None) -> str:
    index = comparison_index_from_label(label)
    if index is None:
        index = fallback_index
    if index is None:
        return "对比文件"
    return f"对比文件{index}"


def merge_reference_candidate(
    candidates: dict[str, ExtractedReference],
    *,
    label: str | None,
    category: str,
    publications: Sequence[str],
    applications: Sequence[str],
    title: str | None,
    context: str,
    source_file: str,
) -> None:
    reference = ExtractedReference(
        label=label,
        category=category,
        publication_numbers=list(publications),
        application_numbers=list(applications),
        title=title,
        context=context,
        source_files=[source_file],
        query_terms=unique_preserve(list(publications) + list(applications) + ([title] if title else [])),
    )
    key = candidate_key(reference)
    if key in candidates:
        candidates[key].merge(reference)
    else:
        candidates[key] = reference


def collect_window(lines: Sequence[str], index: int, radius: int = 2) -> str:
    start = max(0, index - radius)
    end = min(len(lines), index + radius + 1)
    return "\n".join(lines[start:end])


def split_reference_blocks(lines: Sequence[str], citation_only: bool = True) -> list[str]:
    blocks: list[list[str]] = []
    current: list[str] = []

    def flush() -> None:
        nonlocal current
        if not current:
            return
        block_text = "\n".join(current)
        if not citation_only or looks_like_citation_context(block_text):
            blocks.append(current)
        current = []

    for line in lines:
        line_labels = extract_labels(line)
        line_has_identifiers = bool(extract_publication_numbers(line) or extract_application_numbers(line))
        line_is_citation = looks_like_citation_context(line)

        if line_labels:
            flush()
            current = [line]
            continue

        if not line_has_identifiers and not line_is_citation:
            flush()
            continue

        if current:
            current.append(line)
        else:
            current = [line]

    flush()
    return ["\n".join(block) for block in blocks]


def extract_references_from_text(
    text: str,
    source_file: str,
    citation_only: bool = True,
) -> list[ExtractedReference]:
    lines = [normalize_space(line) for line in text.splitlines() if normalize_space(line)]
    candidates: dict[str, ExtractedReference] = {}
    blocks = split_reference_blocks(lines, citation_only=citation_only)
    for block in blocks:
        publications = extract_publication_numbers(block)
        applications = extract_application_numbers(block)
        if not publications and not applications:
            continue
        if citation_only and not looks_like_citation_context(block):
            continue
        labels = extract_labels(block)
        title = extract_title(block)
        if not labels:
            labels = [None]
        for label in labels:
            merge_reference_candidate(
                candidates,
                label=label,
                category=infer_category(label, block),
                publications=publications,
                applications=applications,
                title=title,
                context=block,
                source_file=source_file,
            )
    return list(candidates.values())


def extract_explicit_comparison_references_from_text(
    text: str,
    source_file: str,
) -> list[ExtractedReference]:
    lines = [normalize_space(line) for line in text.splitlines() if normalize_space(line)]
    candidates: dict[str, ExtractedReference] = {}

    for line in lines:
        publications = extract_publication_numbers(line)
        if not publications:
            continue
        index = comparison_index_from_label(line)
        if index is None:
            continue
        merge_reference_candidate(
            candidates,
            label=canonical_comparison_label(line, index),
            category="comparison",
            publications=publications,
            applications=extract_application_numbers(line),
            title=extract_title(line),
            context=line,
            source_file=source_file,
        )

    full_text = "\n".join(lines)
    section_start = full_text.find(COMPARISON_SECTION_MARKER)
    if section_start != -1:
        section = full_text[section_start:]
        end_positions = [
            position
            for marker in COMPARISON_SECTION_END_MARKERS
            if (position := section.find(marker)) != -1
        ]
        if end_positions:
            section = section[: min(end_positions)]
        for line in section.splitlines():
            number_match = re.match(r"^\s*(\d{1,3})\b", line)
            if not number_match:
                continue
            publications = extract_publication_numbers(line)
            if not publications:
                continue
            merge_reference_candidate(
                candidates,
                label=canonical_comparison_label(None, int(number_match.group(1))),
                category="comparison",
                publications=publications,
                applications=extract_application_numbers(line),
                title=extract_title(line),
                context=line,
                source_file=source_file,
            )

    references = list(candidates.values())
    references.sort(
        key=lambda item: (
            comparison_index_from_label(item.label) if comparison_index_from_label(item.label) is not None else 999,
            item.publication_numbers[0] if item.publication_numbers else "",
        )
    )
    return references


def extract_original_application_reference_from_text(
    text: str,
    source_file: str,
) -> ExtractedReference | None:
    lines = [normalize_space(line) for line in text.splitlines() if normalize_space(line)]
    if not lines:
        return None

    header_text = "\n".join(lines[:80])
    application_match = APPLICATION_NUMBER_FIELD_PATTERN.search(header_text)
    publication_match = PUBLICATION_NUMBER_FIELD_PATTERN.search(header_text)
    title_match = INVENTION_TITLE_FIELD_PATTERN.search(header_text)

    if not application_match and not publication_match:
        return None

    applications: list[str] = []
    if application_match:
        application_number = normalize_identifier(application_match.group(1))
        if not re.match(r"^[A-Z]{2}", application_number):
            applications.append(f"CN{application_number}")
        applications.append(application_number)

    publications = [normalize_identifier(publication_match.group(1))] if publication_match else []
    title = normalize_space(title_match.group(1)) if title_match else None
    return ExtractedReference(
        label=SOURCE_APPLICATION_LABEL,
        category="source-application",
        publication_numbers=unique_preserve(publications),
        application_numbers=unique_preserve(applications),
        title=title,
        context=header_text,
        source_files=[source_file],
        query_terms=unique_preserve(publications + applications + ([title] if title else [])),
    )


def extract_original_application_reference(input_files: Sequence[Path]) -> dict | None:
    for path in input_files:
        text = read_document_text(path)
        reference = extract_original_application_reference_from_text(text, str(path))
        if reference:
            return reference.to_dict()
    return None


def extract_references(
    input_files: Sequence[Path],
    citation_only: bool = True,
    comparison_only: bool = False,
) -> dict:
    merged: dict[str, ExtractedReference] = {}
    for path in input_files:
        text = read_document_text(path)
        if comparison_only:
            extracted_references = extract_explicit_comparison_references_from_text(text, str(path))
        else:
            extracted_references = extract_references_from_text(text, str(path), citation_only=citation_only)
        for reference in extracted_references:
            key = candidate_key(reference)
            if key in merged:
                merged[key].merge(reference)
            else:
                merged[key] = reference
    references = sorted(
        (reference.to_dict() for reference in merged.values()),
        key=lambda item: (
            0 if comparison_index_from_label(item["label"]) is not None else {"comparison": 1, "evidence": 2, "other": 3}.get(item["category"], 9),
            comparison_index_from_label(item["label"]) if comparison_index_from_label(item["label"]) is not None else 999,
            item["label"] or "",
            item["publication_numbers"][0] if item["publication_numbers"] else "",
            item["application_numbers"][0] if item["application_numbers"] else "",
        ),
    )
    return {
        "generated_at": now_iso(),
        "citation_only": citation_only,
        "comparison_only": comparison_only,
        "input_files": [str(path) for path in input_files],
        "references": references,
    }


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def requests_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    retry = Retry(
        total=3,
        backoff_factor=0.8,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET", "HEAD"],
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def curl_fetch_text(url: str, timeout: int = 30) -> str:
    result = subprocess.run(
        [
            "curl.exe",
            "-L",
            "-A",
            USER_AGENT,
            "--max-time",
            str(timeout),
            url,
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="ignore",
        check=True,
    )
    return result.stdout


def curl_download(url: str, destination: Path, timeout: int = 60) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            "curl.exe",
            "-L",
            "-A",
            USER_AGENT,
            "--max-time",
            str(timeout),
            url,
            "-o",
            str(destination),
        ],
        check=True,
    )


def parse_patent_page(html: str, url: str) -> dict | None:
    soup = BeautifulSoup(html, "html.parser")
    pdf_meta = soup.find("meta", attrs={"name": "citation_pdf_url"})
    if not pdf_meta or not pdf_meta.get("content"):
        return None
    title_meta = soup.find("meta", attrs={"name": "DC.title"})
    title = normalize_space(title_meta.get("content", "")) if title_meta else None
    page_title = normalize_space(soup.title.text) if soup.title else None
    return {
        "page_url": url,
        "pdf_url": pdf_meta["content"],
        "title": title or page_title,
        "resolver": "direct-page",
    }


def page_contains_application(html: str, applications: Sequence[str]) -> bool:
    upper_html = html.upper()
    for application in applications:
        for variant in application_search_variants(application):
            if variant and variant in upper_html:
                return True
    return False


def fetch_patent_page_by_publication(
    publication_number: str,
    session: requests.Session | None = None,
    timeout: int = 20,
    language: str = "en",
) -> dict | None:
    session = session or requests_session()
    publication = normalize_identifier(publication_number)
    url = f"https://patents.google.com/patent/{publication}/{language}"
    html = None
    final_url = url
    try:
        response = session.get(url, timeout=timeout)
        if response.status_code == 200:
            html = response.text
            final_url = str(response.url)
    except requests.RequestException:
        html = None

    if html is None:
        try:
            html = curl_fetch_text(url, timeout=max(timeout, 30))
        except subprocess.CalledProcessError:
            return None

    parsed = parse_patent_page(html, final_url)
    if not parsed:
        return None
    parsed["publication_number"] = publication
    return parsed


def application_search_variants(application_number: str) -> list[str]:
    normalized = normalize_identifier(application_number)
    variants = [normalized]
    if normalized and not re.match(r"^[A-Z]{2}", normalized):
        variants.append(f"CN{normalized}")
    if "." in normalized:
        variants.append(normalized.split(".", 1)[0])
        if normalized and not re.match(r"^[A-Z]{2}", normalized):
            variants.append(f"CN{normalized.split('.', 1)[0]}")
    digits = re.sub(r"\D", "", normalized)
    if digits:
        variants.append(digits)
        if normalized and not re.match(r"^[A-Z]{2}", normalized):
            variants.append(f"CN{digits}")
    country_digits = re.match(r"^([A-Z]+)(\d+(?:\.\d)?)$", normalized)
    if country_digits:
        country, number = country_digits.groups()
        variants.append(number)
        if "." in number:
            variants.append(number.split(".", 1)[0])
        digit_only = re.sub(r"\D", "", number)
        if len(digit_only) > 1:
            variants.append(country + digit_only[:-1])
            variants.append(digit_only[:-1])
    return unique_preserve(variants)


def title_similarity(a: str | None, b: str | None) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(a=normalize_title(a), b=normalize_title(b)).ratio()


class PlaywrightSearchResolver:
    def __init__(self, headless: bool = True, timeout_ms: int = 20000):
        self.headless = headless
        self.timeout_ms = timeout_ms
        self._playwright = None
        self._browser = None
        self._page = None

    def __enter__(self) -> "PlaywrightSearchResolver":
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        self.close()

    def _ensure_page(self):
        if self._page is not None:
            return self._page
        from playwright.sync_api import sync_playwright

        self._playwright = sync_playwright().start()
        self._browser = self._playwright.chromium.launch(headless=self.headless)
        self._page = self._browser.new_page()
        self._page.set_default_timeout(self.timeout_ms)
        return self._page

    def close(self) -> None:
        if self._page is not None:
            self._page.close()
            self._page = None
        if self._browser is not None:
            self._browser.close()
            self._browser = None
        if self._playwright is not None:
            self._playwright.stop()
            self._playwright = None

    def search(self, query: str, max_results: int = 5) -> list[dict]:
        page = self._ensure_page()
        url = f"https://patents.google.com/?q={quote_plus(query)}"
        last_error = None
        for _ in range(3):
            try:
                page.goto(url, wait_until="load")
                page.wait_for_timeout(3500)
                last_error = None
                break
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                page.wait_for_timeout(1500)
        if last_error is not None:
            raise last_error
        if "No results found." in page.locator("body").inner_text():
            return []
        results: list[dict] = []
        count = page.locator("search-result-item").count()
        for index in range(min(count, max_results)):
            item = page.locator("search-result-item").nth(index)
            state_modifier = item.locator("state-modifier").first
            data_result = state_modifier.get_attribute("data-result")
            pdf_link = item.locator("a.pdfLink")
            pdf_url = pdf_link.first.get_attribute("href") if pdf_link.count() else None
            visible_publication = None
            if pdf_link.count():
                visible_publication = normalize_identifier(pdf_link.first.inner_text())
            title = normalize_space(item.locator("h3").first.inner_text())
            metadata = normalize_space(item.locator("h4.metadata").first.inner_text()) if item.locator("h4.metadata").count() else None
            snippet = normalize_space(item.locator("raw-html").last.inner_text()) if item.locator("raw-html").count() else None
            results.append(
                {
                    "query": query,
                    "title": title,
                    "metadata": metadata,
                    "snippet": snippet,
                    "publication_number": visible_publication,
                    "page_url": f"https://patents.google.com/{data_result}" if data_result else None,
                    "pdf_url": pdf_url,
                    "resolver": "playwright-search",
                }
            )
        return results


def score_search_result(result: dict, reference: dict) -> float:
    score = 0.0
    result_pub = normalize_identifier(result.get("publication_number") or "")
    result_text = " ".join(
        normalize_space(part)
        for part in [result.get("title"), result.get("metadata"), result.get("snippet"), result.get("page_url")]
        if part
    )
    result_text_upper = result_text.upper()
    for publication in reference.get("publication_numbers", []):
        publication_normalized = normalize_identifier(publication)
        if publication_normalized == result_pub:
            score += 100
        elif publication_normalized and publication_normalized in result_text_upper:
            score += 80
    for application in reference.get("application_numbers", []):
        for variant in application_search_variants(application):
            if variant and variant in result_text_upper:
                score += 40
    score += title_similarity(reference.get("title"), result.get("title")) * 30
    if result.get("pdf_url"):
        score += 5
    return score


def strip_kind_code(identifier: str) -> str:
    normalized = normalize_identifier(identifier)
    return re.sub(r"[A-Z]\d?$", "", normalized)


def has_publication_match(result: dict, reference: dict) -> bool:
    result_pub = normalize_identifier(result.get("publication_number") or "")
    result_page = normalize_identifier(result.get("page_url") or "")
    result_text = " ".join(
        normalize_space(part)
        for part in [result.get("title"), result.get("metadata"), result.get("snippet"), result.get("page_url")]
        if part
    ).upper()
    for publication in reference.get("publication_numbers", []):
        normalized = normalize_identifier(publication)
        if not normalized:
            continue
        if normalized == result_pub:
            return True
        if normalized and normalized in result_page:
            return True
        if normalized and normalized in result_text:
            return True
        core = strip_kind_code(normalized)
        result_core = strip_kind_code(result_pub)
        if core and result_core and core == result_core:
            return True
    return False


def has_application_match(result: dict, reference: dict) -> bool:
    result_text = " ".join(
        normalize_space(part)
        for part in [result.get("title"), result.get("metadata"), result.get("snippet"), result.get("page_url")]
        if part
    ).upper()
    for application in reference.get("application_numbers", []):
        for variant in application_search_variants(application):
            if variant and variant in result_text:
                return True
    return False


def is_acceptable_search_result(result: dict, reference: dict) -> bool:
    publication_match = has_publication_match(result, reference)
    application_match = has_application_match(result, reference)
    similarity = title_similarity(reference.get("title"), result.get("title"))

    if publication_match:
        return True
    if application_match and (similarity >= 0.45 or not reference.get("title")):
        return True

    has_number_target = bool(reference.get("publication_numbers") or reference.get("application_numbers"))
    if not has_number_target and reference.get("title") and similarity >= 0.86:
        return True

    if has_number_target and reference.get("title") and similarity >= 0.92:
        return True

    return False


def choose_best_search_result(results: Sequence[dict], reference: dict) -> dict | None:
    if not results:
        return None
    ranked = sorted(results, key=lambda item: score_search_result(item, reference), reverse=True)
    best = ranked[0]
    if not is_acceptable_search_result(best, reference):
        return None
    return best


def resolve_reference(
    reference: dict,
    session: requests.Session | None = None,
    search_resolver: PlaywrightSearchResolver | None = None,
    timeout: int = 20,
) -> dict | None:
    session = session or requests_session()
    for publication in reference.get("publication_numbers", []):
        try:
            resolved = fetch_patent_page_by_publication(publication, session=session, timeout=timeout)
        except requests.RequestException:
            resolved = None
        if resolved:
            resolved["matched_on"] = publication
            return resolved

    queries: list[str] = []
    queries.extend(reference.get("publication_numbers", []))
    for application in reference.get("application_numbers", []):
        queries.extend(application_search_variants(application))
    if reference.get("title"):
        queries.append(reference["title"])
        queries.append(f"\"{reference['title']}\"")

    queries = unique_preserve(normalize_space(query) for query in queries if normalize_space(query))
    if not queries:
        return None

    owns_resolver = False
    if search_resolver is None:
        search_resolver = PlaywrightSearchResolver()
        owns_resolver = True

    try:
        for query in queries:
            results = search_resolver.search(query)
            normalized_query = normalize_identifier(query)
            if (
                reference.get("label") == SOURCE_APPLICATION_LABEL
                and re.fullmatch(r"[A-Z]{2}\d{8,14}(?:\.\d)?", normalized_query)
                and results
                and results[0].get("publication_number")
            ):
                top_result = results[0]
                top_result["matched_on"] = query
                return top_result
            best = choose_best_search_result(results, reference)
            if not best:
                if reference.get("application_numbers"):
                    for candidate in results[:3]:
                        if not candidate.get("page_url"):
                            continue
                        response = session.get(candidate["page_url"], timeout=timeout)
                        if response.status_code != 200:
                            continue
                        if not page_contains_application(response.text, reference["application_numbers"]):
                            continue
                        parsed = parse_patent_page(response.text, candidate["page_url"])
                        if not parsed:
                            continue
                        candidate.update(parsed)
                        candidate["matched_on"] = query
                        candidate["resolver"] = "playwright-search+page-verify"
                        return candidate
                continue
            if best.get("page_url") and not best.get("pdf_url"):
                response = session.get(best["page_url"], timeout=timeout)
                if response.status_code == 200:
                    parsed = parse_patent_page(response.text, best["page_url"])
                    if parsed:
                        best.update(parsed)
            if best.get("pdf_url"):
                best["matched_on"] = query
                return best
    finally:
        if owns_resolver:
            search_resolver.close()
    return None


def download_file(url: str, destination: Path, session: requests.Session | None = None, timeout: int = 30) -> None:
    session = session or requests_session()
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        with session.get(url, timeout=timeout, stream=True) as response:
            response.raise_for_status()
            with destination.open("wb") as file_handle:
                for chunk in response.iter_content(chunk_size=1024 * 128):
                    if chunk:
                        file_handle.write(chunk)
        return
    except requests.RequestException:
        pass

    curl_download(url, destination, timeout=max(timeout, 60))


def file_stem_for_reference(reference: dict, resolved: dict | None = None) -> str:
    parts: list[str] = []
    label = normalize_space(reference.get("label") or "")
    if label:
        parts.append(f"[{label}]")
    identifier = None
    if resolved and resolved.get("publication_number"):
        identifier = resolved["publication_number"]
    elif reference.get("publication_numbers"):
        identifier = reference["publication_numbers"][0]
    elif reference.get("application_numbers"):
        identifier = reference["application_numbers"][0]
    if identifier:
        parts.append(identifier)
    title = None
    if resolved and resolved.get("title"):
        title = resolved["title"]
    elif reference.get("title"):
        title = reference["title"]
    if title:
        parts.append(title)
    return sanitize_filename(" - ".join(parts))


def comparison_file_stem(reference: dict, resolved: dict | None = None, fallback_index: int | None = None) -> str:
    label = canonical_comparison_label(reference.get("label"), fallback_index=fallback_index)
    identifier = None
    if resolved and resolved.get("publication_number"):
        identifier = resolved["publication_number"]
    elif reference.get("publication_numbers"):
        identifier = reference["publication_numbers"][0]
    elif reference.get("application_numbers"):
        identifier = reference["application_numbers"][0]

    stem = label
    if identifier:
        stem = f"{stem}-{identifier}"
    return sanitize_filename(stem)


def source_application_file_stem(reference: dict, resolved: dict | None = None) -> str:
    identifier = None
    if resolved and resolved.get("publication_number"):
        identifier = resolved["publication_number"]
    elif reference.get("publication_numbers"):
        identifier = reference["publication_numbers"][0]
    elif reference.get("application_numbers"):
        identifier = reference["application_numbers"][0]

    stem = SOURCE_APPLICATION_LABEL
    if identifier:
        stem = f"{stem}-{identifier}"
    return sanitize_filename(stem)


def category_output_dir(output_root: Path, category: str) -> Path:
    if category == "comparison":
        return output_root / COMPARISON_DIRNAME
    if category == "evidence":
        return output_root / EVIDENCE_DIRNAME
    return output_root / OTHER_DIRNAME


def infer_project_dir(inputs: Sequence[Path]) -> Path:
    if not inputs:
        return Path.cwd()
    directories = [path if path.is_dir() else path.parent for path in inputs]
    common = Path(os.path.commonpath([str(directory) for directory in directories]))
    return common


def manifest_path(output_root: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return output_root / f"manifest-{timestamp}.json"
