#!/usr/bin/env python3
"""Convert a concise external patent communication Markdown file to DOCX.

This script intentionally implements a fixed style rather than a general
Markdown renderer. It is for short Chinese patent external briefs.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - clear CLI failure path
    raise SystemExit(
        "python-docx is required. Prefer the bundled Codex Python runtime: "
        "/Users/chrynos/.cache/codex-runtimes/codex-primary-runtime/"
        "dependencies/python/bin/python3"
    ) from exc


DEFAULT_FONT = "Microsoft YaHei"
TITLE_COLOR = "0B2545"
HEADING_BLUE = "2E74B5"
META_GRAY = "555555"
BODY_BLACK = "000000"
METADATA_LABELS = {
    "申请号",
    "发明名称",
    "申请人",
    "案号",
    "案件编号",
    "客户",
    "收件人",
    "联系人",
    "发明人",
    "主题",
    "日期",
    "通知书名称",
    "决定号",
    "审查员",
    "对比文件",
    "答复期限",
    "处理期限",
    "结论",
    "建议",
}


def set_run_font(run, *, font=DEFAULT_FONT, size=None, color=None, bold=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_paragraph(paragraph, *, before=0, after=6, line=1.10, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_style_font(doc, style_name, *, size=11, color=BODY_BLACK, bold=False):
    style = doc.styles[style_name]
    style.font.name = DEFAULT_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), DEFAULT_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return style


def add_title_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), HEADING_BLUE)
    p_bdr.append(bottom)


def add_markdown_runs(paragraph, text, *, size=11, color=BODY_BLACK, bold=False):
    """Add text with minimal inline **bold** support."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else part
        run = paragraph.add_run(value)
        set_run_font(run, size=size, color=color, bold=(bold or is_bold))


def is_metadata_line(line):
    if "：" not in line:
        return False
    label, value = line.split("：", 1)
    return label.strip() in METADATA_LABELS and bool(value.strip())


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = set_style_font(doc, "Normal", size=11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = set_style_font(doc, "Title", size=20, color=TITLE_COLOR, bold=True)
    title.paragraph_format.space_after = Pt(10)

    h1 = set_style_font(doc, "Heading 1", size=15, color=HEADING_BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.line_spacing = 1.10

    h2 = set_style_font(doc, "Heading 2", size=13, color=HEADING_BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.line_spacing = 1.10


def add_metadata_paragraph(doc, line):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=2)
    label, value = line.split("：", 1)
    label_run = paragraph.add_run(label.strip() + "：")
    set_run_font(label_run, size=10.5, color=META_GRAY, bold=True)
    value_run = paragraph.add_run(value.strip())
    set_run_font(value_run, size=10.5, color=META_GRAY)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_paragraph(paragraph, after=4, line=1.10)
    add_markdown_runs(paragraph, text, size=11)


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    set_paragraph(paragraph, after=4, line=1.10)
    add_markdown_runs(paragraph, text, size=11)


def markdown_to_docx(markdown_path: Path, output_path: Path, *, author: str | None = None):
    text = markdown_path.read_text(encoding="utf-8-sig")
    doc = Document()
    configure_document(doc)

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            set_paragraph(paragraph, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_title_bottom_rule(paragraph)
            run = paragraph.add_run(line[2:].strip())
            set_run_font(run, size=20, color=TITLE_COLOR, bold=True)
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            set_paragraph(paragraph, before=16, after=8)
            run = paragraph.add_run(line[3:].strip())
            set_run_font(run, size=15, color=HEADING_BLUE, bold=True)
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            set_paragraph(paragraph, before=12, after=6)
            run = paragraph.add_run(line[4:].strip())
            set_run_font(run, size=13, color=HEADING_BLUE, bold=True)
        elif is_metadata_line(line):
            add_metadata_paragraph(doc, line)
        elif re.match(r"^[-*]\s+", line):
            add_bullet(doc, re.sub(r"^[-*]\s+", "", line))
        elif re.match(r"^\d+[.)、]\s+", line):
            add_numbered(doc, re.sub(r"^\d+[.)、]\s+", "", line))
        else:
            paragraph = doc.add_paragraph()
            set_paragraph(paragraph)
            add_markdown_runs(paragraph, line)

    title_text = next((line[2:].strip() for line in text.splitlines() if line.startswith("# ")), "")
    doc.core_properties.title = title_text or "专利对外沟通文件"
    doc.core_properties.subject = "专利对外沟通文件"
    if author:
        doc.core_properties.author = author

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Convert patent external communication Markdown to fixed-layout DOCX.")
    parser.add_argument("markdown", type=Path, help="Input Markdown file")
    parser.add_argument("--output", "-o", type=Path, required=True, help="Output DOCX path")
    parser.add_argument("--author", default=None, help="Optional DOCX core-property author")
    args = parser.parse_args()

    markdown_to_docx(args.markdown, args.output, author=args.author)
    print(args.output)


if __name__ == "__main__":
    main()
