#!/usr/bin/env python3
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


TABLE_HEADERS = [
    "序号",
    "原始交底案号",
    "原始交底名称",
    "方案层级",
    "当前提案名称",
    "申请类型",
    "对应关系及关键技术手段",
]

TABLE_WIDTHS = [0.42, 1.35, 2.20, 0.95, 2.05, 0.70, 2.78]


def require_docx() -> None:
    try:
        import docx  # noqa: F401
    except Exception as exc:  # pragma: no cover
        raise SystemExit(
            "Missing python-docx. Install it or run with the Codex bundled Python runtime."
        ) from exc


def font_run(run, name: str, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None, size: float = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text or "")
    font_run(r, "宋体", size, bold=bold, color=color)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_doc(doc: Document, footer_text: str) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = doc.styles[style_name]
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.bold = True
        st.font.color.rgb = RGBColor(31, 78, 121)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12.5)
    doc.styles["Heading 3"].font.size = Pt(11)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(footer_text)
    font_run(run, "宋体", 9, color="666666")


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    font_run(r, "微软雅黑", 20, bold=True, color="1F4E79")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    font_run(r2, "微软雅黑", 11, color="595959")


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font_run(r, "宋体", 10.5)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    r = p.add_run(text)
    font_run(r, "微软雅黑", 14 if level == 1 else 12, bold=True, color="1F4E79")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        font_run(r, "宋体", 10)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    font_run(r, "宋体", 9)
    r.italic = True


def resolve_path(raw: str, json_dir: Path) -> Path:
    path = Path(raw).expanduser()
    if path.is_absolute():
        return path
    candidate = json_dir / path
    if candidate.exists():
        return candidate
    return Path.cwd() / path


def add_image(doc: Document, image: dict[str, Any], json_dir: Path) -> None:
    raw_path = image.get("path")
    if not raw_path:
        return
    path = resolve_path(str(raw_path), json_dir)
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    width = float(image.get("width_inches", 6.2))
    run.add_picture(str(path), width=Inches(width))
    caption = image.get("caption")
    if caption:
        add_caption(doc, str(caption))


def add_overview_table(doc: Document, cases: list[dict[str, Any]]) -> None:
    add_heading(doc, "一、提案总览", 1)
    table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_row = table.rows[0]
    keep_row_together(header_row)
    repeat_table_header(header_row)
    for idx, header in enumerate(TABLE_HEADERS):
        cell = header_row.cells[idx]
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8)
        cell.width = Inches(TABLE_WIDTHS[idx])

    for case in cases:
        row = table.add_row()
        keep_row_together(row)
        values = [
            str(case.get("no", "")),
            str(case.get("source_no", "")),
            str(case.get("source_name", "")),
            str(case.get("layer", "")),
            str(case.get("title", "")),
            str(case.get("application_type", case.get("type", ""))),
            "对应当前提案。" + str(case.get("table_means", "")),
        ]
        for idx, value in enumerate(values):
            set_cell_text(row.cells[idx], value, size=8)
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if idx == 5:
                set_cell_shading(row.cells[idx], "F8FAFC")
    doc.add_paragraph()


def add_case_details(doc: Document, cases: list[dict[str, Any]], json_dir: Path) -> None:
    add_heading(doc, "二、提案技术内容", 1)
    for case in cases:
        no = case.get("no", "")
        title = case.get("title", "")
        add_heading(doc, f"{no}. {title}", 2)
        add_body(doc, f"对应原始交底：{case.get('source_no', '')}，{case.get('source_name', '')}。")
        if case.get("summary"):
            add_body(doc, str(case["summary"]))

        add_heading(doc, "技术问题", 3)
        add_body(doc, str(case.get("problem", "")))

        add_heading(doc, "关键技术手段", 3)
        add_bullets(doc, [str(item) for item in case.get("means", [])])

        add_heading(doc, "主要技术效果", 3)
        add_body(doc, str(case.get("effect", "")))

        drafting = case.get("drafting") or case.get("preliminary_drafting")
        if drafting:
            add_heading(doc, "初步撰写", 3)
            if isinstance(drafting, list):
                for item in drafting:
                    if item:
                        add_body(doc, str(item))
            else:
                add_body(doc, str(drafting))

        images = case.get("images") or []
        if images:
            add_heading(doc, "图示", 3)
            for image in images:
                add_image(doc, image, json_dir)


def build_docx(data: dict[str, Any], output: Path, json_dir: Path) -> None:
    title = str(data.get("title", "专利提案清单"))
    subtitle = str(data.get("subtitle", "专利布局评估稿"))
    footer = str(data.get("footer", title))
    cases = data.get("cases", [])
    if not isinstance(cases, list) or not cases:
        raise SystemExit("JSON must contain a non-empty 'cases' list.")

    doc = Document()
    configure_doc(doc, footer)
    add_title(doc, title, subtitle)
    for paragraph in data.get("overview", []):
        add_body(doc, str(paragraph))
    add_overview_table(doc, cases)
    add_case_details(doc, cases, json_dir)

    props = doc.core_properties
    props.title = title
    props.subject = "专利提案清单"
    props.author = str(data.get("author", "专利代理师"))
    props.comments = ""
    props.keywords = str(data.get("keywords", "专利提案; 专利布局"))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main(argv: list[str]) -> int:
    require_docx()
    if len(argv) != 3:
        print("Usage: build_proposal_list_docx.py data.json output.docx", file=sys.stderr)
        return 2
    data_path = Path(argv[1]).expanduser().resolve()
    output = Path(argv[2]).expanduser().resolve()
    data = json.loads(data_path.read_text(encoding="utf-8"))
    build_docx(data, output, data_path.parent)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
